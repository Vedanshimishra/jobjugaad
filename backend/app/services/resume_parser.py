"""Resume ingestion: file -> text -> structured profile facts."""

from __future__ import annotations

import base64
import logging
import re
from datetime import date
from pathlib import Path

from pydantic import BaseModel, Field

from app import llm
from app.models import CandidateProfile
from app.services.skills import extract_skills, normalize_skills

log = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


# ---------- schema ----------

class Link(BaseModel):
    label: str
    url: str


class Education(BaseModel):
    school: str
    degree: str = Field(description="e.g. 'B.S.', 'M.S.', 'B.Tech', 'PhD'")
    field: str
    start: str = Field(description="YYYY-MM or YYYY, empty if unknown")
    end: str = Field(description="YYYY-MM or YYYY (expected date if in progress), empty if unknown")
    gpa: str


class Experience(BaseModel):
    company: str
    title: str
    location: str
    start: str = Field(description="YYYY-MM, empty if unknown")
    end: str = Field(description="YYYY-MM or 'present'")
    is_internship: bool
    bullets: list[str]
    skills: list[str]


class Project(BaseModel):
    name: str
    description: str
    bullets: list[str]
    skills: list[str]
    url: str


class ParsedResume(BaseModel):
    full_name: str
    email: str
    phone: str
    location: str
    links: list[Link]
    headline: str = Field(description="One-line professional headline inferred from the resume")
    summary: str
    education: list[Education]
    experience: list[Experience]
    projects: list[Project]
    skills: list[str]
    certifications: list[str]
    graduation_year: int | None = Field(description="Year of the most recent (or expected) degree completion")


SYSTEM = """You extract structured data from resumes for a job-search assistant.
Rules:
- Copy facts exactly as written. Never invent employers, dates, metrics, or skills.
- Use empty strings / empty lists for anything absent.
- Normalize dates to YYYY-MM (or YYYY if only the year is given); use 'present' for current roles.
- 'skills' is the de-duplicated union of explicitly listed skills and technologies clearly used in experience/projects.
- Bullets should be the original bullet text, lightly cleaned of formatting artifacts."""


# ---------- text extraction ----------

def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join((p.extract_text() or "") for p in reader.pages).strip()
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts).strip()
    if ext in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore").strip()
    raise ValueError(f"Unsupported resume format: {ext}")


# ---------- parsing ----------

def parse_with_llm(path: Path, text: str) -> ParsedResume:
    if path.suffix.lower() == ".pdf" and path.stat().st_size < 20 * 1024 * 1024:
        data = base64.standard_b64encode(path.read_bytes()).decode()
        content = [
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": data}},
            {"type": "text", "text": "Extract this resume into the schema."},
        ]
    else:
        content = f"<resume>\n{text}\n</resume>\n\nExtract this resume into the schema."
    return llm.structured(system=SYSTEM, content=content, output=ParsedResume, effort="medium")


_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE = re.compile(r"(\+?\d[\d\s().-]{8,}\d)")
_URL = re.compile(r"(https?://\S+|(?:www\.)?(?:linkedin\.com|github\.com)/\S+)", re.I)


def parse_heuristic(text: str) -> ParsedResume:
    """Best-effort fallback when the LLM is unavailable. The user can correct the result."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    email = (_EMAIL.search(text) or [""])[0] if _EMAIL.search(text) else ""
    phone_m = _PHONE.search(text)
    links = [Link(label="github" if "github" in u.lower() else "linkedin" if "linkedin" in u.lower() else "link", url=u.rstrip(".,;)"))
             for u in dict.fromkeys(_URL.findall(text))]
    years = [int(y) for y in re.findall(r"\b(20[0-4]\d)\b", text)]
    grad = None
    edu_idx = next((i for i, l in enumerate(lines) if re.match(r"(?i)^education\b", l)), None)
    if edu_idx is not None:
        window = " ".join(lines[edu_idx: edu_idx + 8])
        ys = [int(y) for y in re.findall(r"\b(20[0-4]\d)\b", window)]
        grad = max(ys) if ys else None
    elif years:
        grad = max(y for y in years if y <= date.today().year + 5)
    return ParsedResume(
        full_name=lines[0] if lines and len(lines[0]) < 60 else "",
        email=email,
        phone=phone_m.group(1).strip() if phone_m else "",
        location="",
        links=links,
        headline="",
        summary="",
        education=[],
        experience=[],
        projects=[],
        skills=extract_skills(text),
        certifications=[],
        graduation_year=grad,
    )


def parse_resume(path: Path) -> tuple[str, ParsedResume, str]:
    text = extract_text(path)
    try:
        return text, parse_with_llm(path, text), "llm"
    except llm.LLMError as e:
        log.warning("LLM resume parsing unavailable (%s); using heuristic parser", e)
        if not text:
            raise ValueError("Could not extract text from this file (is it a scanned image?).") from e
        return text, parse_heuristic(text), "heuristic"


# ---------- profile building ----------

def _ym(s: str) -> tuple[int, int] | None:
    s = (s or "").strip().lower()
    if s in ("present", "current", "now"):
        t = date.today()
        return t.year, t.month
    m = re.match(r"(\d{4})(?:-(\d{1,2}))?", s)
    if not m:
        return None
    return int(m.group(1)), int(m.group(2) or 6)


def experience_months(experience: list[dict]) -> int:
    """Months of experience, merging overlaps. Internships count at half weight,
    since postings asking for 'N years' usually mean full-time experience."""
    spans: dict[bool, list[tuple[int, int]]] = {True: [], False: []}
    for e in experience:
        a, b = _ym(e.get("start", "")), _ym(e.get("end", "") or "present")
        if not a or not b:
            continue
        s, t = a[0] * 12 + a[1], b[0] * 12 + b[1]
        if t >= s:
            spans[bool(e.get("is_internship"))].append((s, t))

    def merged(ranges: list[tuple[int, int]]) -> int:
        total, cur = 0, None
        for s, t in sorted(ranges):
            if cur and s <= cur[1]:
                cur = (cur[0], max(cur[1], t))
            else:
                if cur:
                    total += cur[1] - cur[0] + 1
                cur = (s, t)
        return total + (cur[1] - cur[0] + 1 if cur else 0)

    return merged(spans[False]) + merged(spans[True]) // 2


def apply_to_profile(profile: CandidateProfile, parsed: ParsedResume) -> None:
    """Copy resume facts onto the profile. Preferences are never touched."""
    p = parsed.model_dump()
    for key in ("full_name", "email", "phone", "location", "headline", "summary"):
        if p[key]:
            setattr(profile, key, p[key])
    profile.links = {l["label"] or f"link{i}": l["url"] for i, l in enumerate(p["links"])}
    profile.education = p["education"]
    profile.experience = p["experience"]
    profile.projects = p["projects"]
    profile.skills = normalize_skills(p["skills"])
    profile.certifications = p["certifications"]
    if p["graduation_year"]:
        profile.graduation_year = p["graduation_year"]
    profile.experience_months = experience_months(p["experience"])
    profile.version = (profile.version or 0) + 1
